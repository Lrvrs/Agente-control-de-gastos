"""Carga de los datos de partida y de los ficheros que sube el alumno."""

import csv
import io
from pathlib import Path
from typing import List

from openpyxl import load_workbook

from dominio.gasto import ConjuntoGastos, Gasto
from dominio.politica import Politica


class ErrorFicheroGastos(Exception):
    """
    Problema en el fichero de gastos que sube el alumno.

    Se define aparte para que la interfaz pueda distinguir un fichero mal
    formado -que el alumno puede corregir- de un fallo interno de la
    aplicacion, y mostrar en cada caso un mensaje que sirva de algo.
    """


class RepositorioDatos:
    """
    Acceso de solo lectura a los ficheros de datos incluidos en el repositorio.

    Se aisla aqui para que el resto de la aplicacion no sepa nada de rutas ni de
    formatos: si manana los gastos vinieran de una hoja de calculo o de una base
    de datos, solo cambiaria esta clase.
    """

    # Nombres de los ficheros de partida. Se declaran como constantes de clase
    # para que la ruta no quede dispersa en varias cadenas literales.
    FICHERO_GASTOS = "gastos_demo.csv"
    FICHERO_POLITICA = "politica_viajes.md"

    # Columnas que debe traer cualquier fichero de gastos. Si falta alguna, el
    # modelo recibiria informacion incompleta y emitiria veredictos sin base.
    COLUMNAS_REQUERIDAS = (
        "id", "empleado", "fecha", "categoria", "ciudad",
        "importe", "moneda", "descripcion", "justificante",
    )

    # Tope de filas admitidas. Protege dos cosas a la vez: el presupuesto de
    # tokens de la llamada al modelo, que es unica para todo el fichero, y la
    # memoria compartida del contenedor gratuito de Streamlit.
    MAXIMO_FILAS = 40

    def __init__(self, carpeta_datos: Path | None = None) -> None:
        """
        Fija la carpeta de datos.

        Por defecto se calcula relativa a este fichero y no al directorio de
        trabajo, porque Streamlit Community Cloud no garantiza cual es el
        directorio actual al arrancar la aplicacion.
        """
        if carpeta_datos is None:
            # padre de infraestructura/ -> raiz del proyecto -> datos/
            carpeta_datos = Path(__file__).resolve().parent.parent / "datos"
        self._carpeta = carpeta_datos

    def cargar_gastos(self) -> ConjuntoGastos:
        """Lee el CSV de gastos de ejemplo y devuelve el conjunto de dominio."""
        ruta = self._carpeta / self.FICHERO_GASTOS

        # Se abre con utf-8-sig para tolerar el marcador de orden de bytes que
        # Excel anade al exportar a CSV en Windows, un origen habitual de fallos.
        with ruta.open("r", encoding="utf-8-sig", newline="") as fichero:
            lector = csv.DictReader(fichero)
            gastos = [self._fila_a_gasto(fila) for fila in lector]

        return ConjuntoGastos(gastos)

    def cargar_gastos_desde_texto(self, contenido: str) -> ConjuntoGastos:
        """
        Construye el conjunto de gastos a partir del contenido de un CSV subido.

        Se separa de `cargar_gastos` porque el fichero que sube un alumno no
        esta en disco: llega en memoria desde el navegador.
        """
        # `splitlines` normaliza los tres estilos de salto de linea posibles,
        # de modo que da igual si el CSV se genero en Windows, macOS o Linux.
        lector = csv.DictReader(contenido.splitlines())
        gastos = [self._fila_a_gasto(fila) for fila in lector]
        return ConjuntoGastos(gastos)

    def cargar_politica(self) -> Politica:
        """Lee la politica de viajes por defecto."""
        ruta = self._carpeta / self.FICHERO_POLITICA
        return Politica(texto=ruta.read_text(encoding="utf-8"))

    def cargar_gastos_subidos(self, nombre: str, contenido: bytes) -> ConjuntoGastos:
        """
        Construye el conjunto de gastos a partir del fichero que sube el alumno.

        Actua como unico punto de entrada para los ficheros externos: decide el
        formato por la extension, delega en el lector correspondiente y aplica
        despues las mismas validaciones a todos. De ese modo un CSV y un Excel
        reciben exactamente el mismo trato.
        """
        # La extension se normaliza porque el sistema de ficheros de quien sube
        # el fichero puede haberla dejado en mayusculas.
        nombre_normalizado = (nombre or "").lower()

        if nombre_normalizado.endswith(".csv"):
            filas = self._leer_filas_csv(contenido)
        elif nombre_normalizado.endswith((".xlsx", ".xlsm")):
            filas = self._leer_filas_excel(contenido)
        else:
            raise ErrorFicheroGastos(
                "Formato no admitido. Sube un fichero .csv o .xlsx."
            )

        # Validaciones comunes a ambos formatos.
        self._validar_filas(filas)

        return ConjuntoGastos([self._fila_a_gasto(fila) for fila in filas])

    def _leer_filas_csv(self, contenido: bytes) -> List[dict]:
        """Decodifica un CSV subido y devuelve sus filas como diccionarios."""
        # Se intentan dos codificaciones antes de rendirse. UTF-8 es lo habitual,
        # pero un CSV exportado desde Excel en Windows suele venir en latin-1 y
        # fallaria con un error incomprensible para el alumno.
        for codificacion in ("utf-8-sig", "latin-1"):
            try:
                texto = contenido.decode(codificacion)
                break
            except UnicodeDecodeError:
                continue
        else:
            raise ErrorFicheroGastos("No se ha podido leer el fichero.")

        # El delimitador tambien varia: una hoja de calculo en espanol exporta
        # con punto y coma. Se detecta contando cual aparece mas en la cabecera.
        primera_linea = texto.splitlines()[0] if texto.splitlines() else ""
        delimitador = ";" if primera_linea.count(";") > primera_linea.count(",") else ","

        lector = csv.DictReader(texto.splitlines(), delimiter=delimitador)
        return [fila for fila in lector]

    def _leer_filas_excel(self, contenido: bytes) -> List[dict]:
        """Lee la primera hoja de un libro de Excel y devuelve sus filas."""
        try:
            # read_only acelera la lectura y data_only devuelve el resultado de
            # las formulas en lugar de su texto, que es lo que interesa aqui.
            libro = load_workbook(
                io.BytesIO(contenido), read_only=True, data_only=True
            )
        except Exception as error:
            raise ErrorFicheroGastos(
                f"No se ha podido abrir el libro de Excel: {error}"
            ) from error

        hoja = libro.active
        filas_brutas = list(hoja.iter_rows(values_only=True))

        # Un libro sin cabecera y al menos una fila de datos no sirve de nada.
        if len(filas_brutas) < 2:
            raise ErrorFicheroGastos("La hoja no contiene datos.")

        # La primera fila son los nombres de columna, normalizados a minusculas
        # y sin espacios para tolerar cabeceras escritas de cualquier manera.
        cabecera = [str(celda or "").strip().lower() for celda in filas_brutas[0]]

        filas: List[dict] = []
        for fila_bruta in filas_brutas[1:]:
            # Se descartan las filas completamente vacias, que son frecuentes al
            # final de una hoja donde alguien ha borrado contenido.
            if all(celda is None or str(celda).strip() == "" for celda in fila_bruta):
                continue

            # zip corta por la lista mas corta, de modo que una fila con menos
            # celdas que la cabecera no provoca un error de indice.
            filas.append(
                {
                    nombre: ("" if valor is None else str(valor).strip())
                    for nombre, valor in zip(cabecera, fila_bruta)
                }
            )

        return filas

    def _validar_filas(self, filas: List[dict]) -> None:
        """Comprueba que las filas leidas sirven para evaluar."""
        # Fichero sin ninguna fila util.
        if not filas:
            raise ErrorFicheroGastos("El fichero no contiene ninguna fila de datos.")

        # Fichero demasiado grande para una sola llamada al modelo.
        if len(filas) > self.MAXIMO_FILAS:
            raise ErrorFicheroGastos(
                f"El fichero tiene {len(filas)} filas y el maximo es "
                f"{self.MAXIMO_FILAS}. Recorta el fichero y vuelve a subirlo."
            )

        # Columnas ausentes. Se nombran todas las que faltan de una vez, en
        # lugar de obligar al alumno a descubrirlas de una en una.
        presentes = {clave.strip().lower() for clave in filas[0].keys() if clave}
        ausentes = [c for c in self.COLUMNAS_REQUERIDAS if c not in presentes]
        if ausentes:
            raise ErrorFicheroGastos(
                "Al fichero le faltan estas columnas: " + ", ".join(ausentes)
            )

    def _fila_a_gasto(self, fila: dict) -> Gasto:
        """
        Convierte una fila del CSV en una entidad Gasto.

        Cada campo se limpia de espacios sobrantes porque un CSV editado a mano
        en una hoja de calculo suele traerlos y provocarian comparaciones
        fallidas mas adelante.
        """
        # El importe se convierte admitiendo la coma decimal, que es lo que
        # escribe una hoja de calculo configurada en espanol.
        importe_bruto = str(fila.get("importe", "0")).strip().replace(",", ".")
        try:
            importe = float(importe_bruto)
        except ValueError:
            # Un importe ilegible se registra como cero: el modelo lo vera y
            # deberia emitir REVISION, que es el comportamiento deseado.
            importe = 0.0

        return Gasto(
            identificador=str(fila.get("id", "")).strip(),
            empleado=str(fila.get("empleado", "")).strip(),
            fecha=str(fila.get("fecha", "")).strip(),
            categoria=str(fila.get("categoria", "")).strip(),
            ciudad=str(fila.get("ciudad", "")).strip(),
            importe=importe,
            moneda=str(fila.get("moneda", "EUR")).strip().upper(),
            descripcion=str(fila.get("descripcion", "")).strip(),
            justificante=str(fila.get("justificante", "N")).strip(),
        )
